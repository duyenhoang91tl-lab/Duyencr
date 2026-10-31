from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
F='Arial'
def font(**k): return Font(name=F, size=k.pop('size',10), **k)
HDR=PatternFill('solid',fgColor='1F3A5F'); INP=PatternFill('solid',fgColor='FFF9C4'); EX=PatternFill('solid',fgColor='EEF2F7')
thin=Side(style='thin',color='BFBFBF'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
MONEY='#,##0;[Red]-#,##0;-'
wb=Workbook()
def head(ws,row,cols,widths):
    for i,(c,w) in enumerate(zip(cols,widths),1):
        x=ws.cell(row=row,column=i,value=c); x.font=font(bold=True,color='FFFFFF'); x.fill=HDR; x.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); x.border=BOX
        ws.column_dimensions[x.column_letter].width=w
    ws.row_dimensions[row].height=30
def title(ws,t,sub):
    ws['A1']=t; ws['A1'].font=font(bold=True,size=14,color='1F3A5F'); ws['A2']=sub; ws['A2'].font=font(italic=True,color='666666'); ws['A2'].alignment=Alignment(wrap_text=False)

# ---- Quy định ----
q=wb.active; q.title='Quy dinh'
title(q,'QUY ĐỊNH THƯỞNG PHẠT — ĐỘI SALE ONLINE + VĂN PHÒNG (phòng kinh doanh)','Riêng, KHÔNG cộng vào các chương trình thưởng của công ty. Ô nền vàng = ô để sửa; các sheet khác tự lấy số tiền từ đây.')
q['A4']='THƯỞNG'; q['A4'].font=font(bold=True,size=11)
rows=[('Số đơn tối thiểu để được thưởng ngày (đơn)',6,'Min6'),('Thưởng sale nhiều đơn nhất trong ngày (đ)',50000,'Day'),('Thưởng top tỷ lệ chốt cao nhất TUẦN (đ)',50000,'Week'),('Thưởng top tỷ lệ chốt cao nhất THÁNG (đ)',100000,'Month')]
for i,(l,v,_) in enumerate(rows,5):
    q.cell(row=i,column=1,value=l).font=font(); c=q.cell(row=i,column=2,value=v); c.font=font(color='0000FF',bold=True); c.fill=INP; c.number_format='#,##0'; c.border=BOX
q['A9']='Đơn tính thưởng = ĐƠN GỐC BASE (theo ngày tạo), không tính theo Pos. Không tính đơn đổi trả, bảo hành, thay dây, huỷ/hoàn.'; q['A9'].font=font(italic=True,color='666666')
q['A10']='Thưởng tuần/tháng: chỉ xét sale có số tương tác trong kỳ ≥ tương tác trung bình của đội. Tỷ lệ chốt = đơn Base ÷ tương tác Pancake. Số đơn tính theo NGƯỜI TẠO ĐƠN (không chia theo Sale bán).'; q['A10'].font=font(italic=True,color='666666')
q['A12']='PHẠT'; q['A12'].font=font(bold=True,size=11)
for i,h in enumerate(['Mã lỗi','Nội dung vi phạm','Số tiền phạt (đ)'],1):
    c=q.cell(row=13,column=i,value=h); c.font=font(bold=True,color='FFFFFF'); c.fill=HDR; c.border=BOX; c.alignment=Alignment(horizontal='center')
pen=[(1,'Tin trôi quá 30 phút không phản hồi (không phải tin nhảy từ Sale khác)',50000),
(2,'Không gửi báo cáo đúng khung giờ (sáng trước 9h30, chiều 13h30, kết thúc ca, tối 20h nếu bật tin)',50000),
(3,'Đã đọc nhưng không phản hồi',50000),
(4,'Đơn gộp Sale bán chung trong 72 giờ: Sale lên đơn không chia doanh thu cho các Sale cùng tham gia',200000),
(5,'Nhắn tin khai thác mong cầu trong 15 phút nhưng không gửi kèm từ 04 hình ảnh hoặc ghi âm',50000)]
for i,(m,n,a) in enumerate(pen,14):
    q.cell(row=i,column=1,value=m).font=font(); q.cell(row=i,column=1).alignment=Alignment(horizontal='center'); q.cell(row=i,column=1).border=BOX
    c=q.cell(row=i,column=2,value=n); c.font=font(); c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=BOX
    c=q.cell(row=i,column=3,value=a); c.font=font(color='0000FF',bold=True); c.fill=INP; c.number_format='#,##0'; c.border=BOX
    q.row_dimensions[i].height=30
q.column_dimensions['A'].width=46; q.column_dimensions['B'].width=88; q.column_dimensions['C'].width=18
q['A20']='Chú thích màu: chữ xanh dương = số nhập tay / quy định · chữ đen = công thức tự tính · nền vàng = ô cần nhập.'; q['A20'].font=font(italic=True,color='666666')

# ---- Danh sach sale ----
d=wb.create_sheet('Danh sach sale')
title(d,'DANH SÁCH SALE ĐỘI ONLINE + VĂN PHÒNG','Điền tên sale (đúng như trên CRM). Sheet "Tong hop thang" và danh sách chọn ở các sheet khác lấy từ đây.')
head(d,4,['Tên sale','Nhóm'],[32,16])
ex=[('Sale A (ví dụ — thay bằng tên thật)','Online'),('Sale B (ví dụ — thay bằng tên thật)','Văn phòng')]
for i in range(5,5+40):
    for j in (1,2):
        c=d.cell(row=i,column=j); c.fill=INP; c.border=BOX; c.font=font(color='0000FF')
for i,(a,b) in enumerate(ex,5): d.cell(row=i,column=1,value=a); d.cell(row=i,column=2,value=b)
SALES="'Danh sach sale'!$A$5:$A$44"

# ---- Nhat ky phat ----
p=wb.create_sheet('Nhat ky phat')
title(p,'NHẬT KÝ PHẠT','Mỗi lần vi phạm ghi 1 dòng: Ngày · Sale · Mã lỗi (1-5). Số tiền tự lấy từ "Quy dinh" — muốn phạt khác mức mặc định thì gõ đè số vào cột "Số tiền".')
head(p,4,['Ngày vi phạm','Sale','Mã lỗi (1-5)','Nội dung vi phạm','Số tiền phạt (đ)','Ghi chú (khách/tin nhắn…)','Người ghi','Tháng'],[14,30,12,52,16,40,16,10])
N0,N1=5,504
for r in range(N0,N1+1):
    for col in (1,2,3,6,7): 
        c=p.cell(row=r,column=col); c.fill=INP; c.border=BOX; c.font=font(color='0000FF')
    p.cell(row=r,column=1).number_format='dd/mm/yyyy'
    d_=p.cell(row=r,column=4,value=f'=IF(C{r}="","",IFERROR(INDEX(\'Quy dinh\'!$B$14:$B$18,MATCH(C{r},\'Quy dinh\'!$A$14:$A$18,0)),"Mã lỗi không hợp lệ"))'); d_.font=font(); d_.border=BOX; d_.alignment=Alignment(wrap_text=True,vertical='top')
    e=p.cell(row=r,column=5,value=f'=IF(C{r}="","",IFERROR(INDEX(\'Quy dinh\'!$C$14:$C$18,MATCH(C{r},\'Quy dinh\'!$A$14:$A$18,0)),0))'); e.font=font(); e.number_format=MONEY; e.border=BOX
    h=p.cell(row=r,column=8,value=f'=IF(A{r}="","",TEXT(A{r},"yyyy-mm"))'); h.font=font(); h.border=BOX; h.alignment=Alignment(horizontal='center')
# vi du
import datetime
p['A5']=datetime.date(2026,10,5); p['B5']='Sale A (ví dụ — thay bằng tên thật)'; p['C5']=4; p['F5']='VÍ DỤ: xoá dòng này khi bắt đầu nhập thật'; p['G5']='Admin'
for col in (1,2,3,6,7): p.cell(row=5,column=col).font=font(color='7F7F7F',italic=True)
dv=DataValidation(type='list',formula1=f'={SALES}',allow_blank=True); dv.error='Chọn tên trong "Danh sach sale"'; p.add_data_validation(dv); dv.add(f'B{N0}:B{N1}')
dv2=DataValidation(type='whole',operator='between',formula1='1',formula2='5',allow_blank=True); dv2.error='Mã lỗi từ 1 đến 5'; p.add_data_validation(dv2); dv2.add(f'C{N0}:C{N1}')
p.freeze_panes='A5'

# ---- Thuong ngay ----
t=wb.create_sheet('Thuong ngay')
title(t,'THƯỞNG NGÀY — SALE NHIỀU ĐƠN NHẤT TRONG NGÀY','Mỗi ngày ghi sale có nhiều đơn Base nhất (đếm theo NGƯỜI TẠO ĐƠN) + số đơn đó (không tính đổi trả, bảo hành, thay dây). Ngày có nhiều người bằng nhau → ghi mỗi người 1 dòng. Đủ tối thiểu thì tự ra tiền thưởng.')
head(t,4,['Ngày','Sale đạt cao nhất','Số đơn Base','Đạt tối thiểu?','Tiền thưởng (đ)','Ghi chú','Tháng'],[14,30,14,16,16,40,10])
T0,T1=5,404
for r in range(T0,T1+1):
    for col in (1,2,3,6):
        c=t.cell(row=r,column=col); c.fill=INP; c.border=BOX; c.font=font(color='0000FF')
    t.cell(row=r,column=1).number_format='dd/mm/yyyy'
    c=t.cell(row=r,column=4,value=f'=IF(C{r}="","",IF(C{r}>=\'Quy dinh\'!$B$5,"Đạt","Chưa đủ"))'); c.font=font(); c.border=BOX; c.alignment=Alignment(horizontal='center')
    c=t.cell(row=r,column=5,value=f'=IF(C{r}="","",IF(AND(B{r}<>"",C{r}>=\'Quy dinh\'!$B$5),\'Quy dinh\'!$B$6,0))'); c.font=font(); c.number_format=MONEY; c.border=BOX
    c=t.cell(row=r,column=7,value=f'=IF(A{r}="","",TEXT(A{r},"yyyy-mm"))'); c.font=font(); c.border=BOX; c.alignment=Alignment(horizontal='center')
t['A5']=datetime.date(2026,10,1); t['B5']='Sale A (ví dụ — thay bằng tên thật)'; t['C5']=7; t['F5']='VÍ DỤ: xoá dòng này khi bắt đầu nhập thật'
for col in (1,2,3,6): t.cell(row=5,column=col).font=font(color='7F7F7F',italic=True)
dv3=DataValidation(type='list',formula1=f'={SALES}',allow_blank=True); t.add_data_validation(dv3); dv3.add(f'B{T0}:B{T1}')
t.freeze_panes='A5'

# ---- Thuong tuan thang ----
w=wb.create_sheet('Thuong tuan-thang')
title(w,'THƯỞNG TỶ LỆ CHỐT CAO NHẤT — TUẦN & THÁNG','Mỗi kỳ liệt kê các sale: tương tác + đơn Base → tỷ lệ chốt tự tính. Nhập "TB tương tác đội" của kỳ; ai đủ điều kiện và đứng đầu thì gõ "x" ở cột "Top" để ra tiền thưởng.')
head(w,4,['Loại kỳ (Tuần/Tháng)','Kỳ (vd 12/10-18/10)','Tháng (yyyy-mm)','Sale','Tương tác (tin nhắn)','Đơn Base','Tỷ lệ chốt','TB tương tác đội','Đủ điều kiện?','Top? (gõ x)','Tiền thưởng (đ)'],[16,22,14,30,14,12,12,14,14,12,16])
W0,W1=5,304
for r in range(W0,W1+1):
    for col in (1,2,3,4,5,6,8,10):
        c=w.cell(row=r,column=col); c.fill=INP; c.border=BOX; c.font=font(color='0000FF')
    c=w.cell(row=r,column=7,value=f'=IF(OR(E{r}="",F{r}=""),"",IF(E{r}=0,0,F{r}/E{r}))'); c.font=font(); c.number_format='0.0%'; c.border=BOX
    c=w.cell(row=r,column=9,value=f'=IF(OR(E{r}="",H{r}=""),"",IF(E{r}>=H{r},"Đủ","Không đủ"))'); c.font=font(); c.border=BOX; c.alignment=Alignment(horizontal='center')
    c=w.cell(row=r,column=11,value=f'=IF(J{r}="","",IF(AND(LOWER(J{r})="x",I{r}="Đủ"),IF(A{r}="Tháng",\'Quy dinh\'!$B$8,IF(A{r}="Tuần",\'Quy dinh\'!$B$7,0)),0))'); c.font=font(); c.number_format=MONEY; c.border=BOX
vals=['Tuần','12/10-18/10','2026-10','Sale A (ví dụ — thay bằng tên thật)',200,30,None,113,None,'x']
for col,v in zip(range(1,11),vals):
    if v is not None: w.cell(row=5,column=col,value=v)
for col in (1,2,3,4,5,6,8,10): w.cell(row=5,column=col).font=font(color='7F7F7F',italic=True)
w['L5']='← VÍ DỤ: xoá dòng này khi bắt đầu nhập thật'; w['L5'].font=font(italic=True,color='7F7F7F')
dv4=DataValidation(type='list',formula1='"Tuần,Tháng"',allow_blank=True); w.add_data_validation(dv4); dv4.add(f'A{W0}:A{W1}')
dv5=DataValidation(type='list',formula1=f'={SALES}',allow_blank=True); w.add_data_validation(dv5); dv5.add(f'D{W0}:D{W1}')
w.freeze_panes='A5'

# ---- Tong hop ----
s=wb.create_sheet('Tong hop thang',0)
title(s,'TỔNG HỢP THƯỞNG PHẠT THÁNG — ĐỘI SALE ONLINE + VĂN PHÒNG','Chọn tháng ở ô B4 (dạng yyyy-mm). Bảng tự cộng từ 3 sheet: Thuong ngay, Thuong tuan-thang, Nhat ky phat. Tên sale lấy từ "Danh sach sale".')
s['A4']='Tháng cần tính (yyyy-mm):'; s['A4'].font=font(bold=True)
s['B4']='2026-10'; s['B4'].font=font(color='0000FF',bold=True); s['B4'].fill=INP; s['B4'].border=BOX; s['B4'].number_format='@'
head(s,6,['Sale','Thưởng ngày (đ)','Thưởng tuần/tháng (đ)','Tổng thưởng (đ)','Phạt (đ)','Số lần phạt','THỰC NHẬN (đ)'],[34,18,22,18,16,12,18])
R0=7
for i in range(40):
    r=R0+i
    a=s.cell(row=r,column=1,value=f'=IF(\'Danh sach sale\'!A{5+i}="","",\'Danh sach sale\'!A{5+i})'); a.font=font(); a.border=BOX
    f={2:f'=IF($A{r}="","",SUMIFS(\'Thuong ngay\'!$E$5:$E$404,\'Thuong ngay\'!$B$5:$B$404,$A{r},\'Thuong ngay\'!$G$5:$G$404,$B$4))',
       3:f'=IF($A{r}="","",SUMIFS(\'Thuong tuan-thang\'!$K$5:$K$304,\'Thuong tuan-thang\'!$D$5:$D$304,$A{r},\'Thuong tuan-thang\'!$C$5:$C$304,$B$4))',
       4:f'=IF($A{r}="","",B{r}+C{r})',
       5:f'=IF($A{r}="","",SUMIFS(\'Nhat ky phat\'!$E$5:$E$504,\'Nhat ky phat\'!$B$5:$B$504,$A{r},\'Nhat ky phat\'!$H$5:$H$504,$B$4))',
       6:f'=IF($A{r}="","",COUNTIFS(\'Nhat ky phat\'!$B$5:$B$504,$A{r},\'Nhat ky phat\'!$H$5:$H$504,$B$4))',
       7:f'=IF($A{r}="","",D{r}-E{r})'}
    for col,fm in f.items():
        c=s.cell(row=r,column=col,value=fm); c.font=font(bold=(col==7)); c.border=BOX; c.number_format=('0' if col==6 else MONEY)
rt=R0+40
s.cell(row=rt,column=1,value='TỔNG').font=font(bold=True)
for col in (2,3,4,5,6,7):
    L=s.cell(row=6,column=col).column_letter
    c=s.cell(row=rt,column=col,value=f'=SUM({L}{R0}:{L}{rt-1})'); c.font=font(bold=True); c.number_format=('0' if col==6 else MONEY); c.border=Border(top=Side(style='medium'))
s.cell(row=rt+2,column=1,value='Ghi chú: "Thuong ngay" chỉ tính dòng có sale + số đơn ≥ mức tối thiểu; "Thuong tuan-thang" chỉ tính dòng gõ "x" và đủ điều kiện tương tác. Nhật ký phạt cộng theo tháng của ngày vi phạm.').font=font(italic=True,color='666666')
s.freeze_panes='A7'
for ws in wb.worksheets:
    ws.sheet_view.showGridLines=False
wb.save('Theo_doi_thuong_phat_doi_Sale.xlsx')
